"""One-off migration of BabyPerfMGT.xlsx (APP_SPEC §7).

Reads Constance's month per Excel row, maps each row to a catalogue skill (MAPPING below,
reviewable in catalog/seed/excel-mapping.tsv) and writes catalog/seed/constance.json:
observations dated birth date + m months, flagged as approximate.
When several Excel rows map to the same skill, the earliest month is kept.
Run: python3 scripts/import_excel.py   (needs: pip install openpyxl)
Then re-run scripts/fill_typical_ages.py to add the typical-age skills again.
"""
import csv
import json
from datetime import date, timedelta
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parent.parent
BIRTH = date(2022, 2, 4)

MAPPING = {
    3: "com.calms_voice", 4: "think.watches_move", 5: "com.loud_sounds", 6: "mob.head_tummy",
    7: "soc.smiles_back", 8: "com.coos", 9: "com.turns_to_voice", 10: "mob.head_tummy",
    11: "hand.holds_toy", 12: "hand.reach_grab", 13: "mob.roll_tummy_back", 14: "soc.laughs",
    15: "mob.head_steady", 16: "hand.mouths_objects", 17: "mob.sit_leaning", 18: "mob.straight_arms",
    19: "mob.sit_alone", 20: "com.babbles", 21: "hand.transfer", 22: "mob.legs_push",
    23: "hand.bangs_objects", 24: "com.responds_name", 25: "soc.peekaboo", 26: "soc.stranger_anxiety",
    27: "mob.crawl", 28: "mob.stand_assisted", 29: "com.understands_no", 30: "soc.pat_a_cake",
    31: "com.waves", 32: "com.directions_gesture", 33: "hand.pincer", 34: "lit.looks_pages",
    35: "mob.walk_alone", 36: "com.first_words", 37: "com.points_ask", 38: "auto.finger_feeds",
    39: "auto.cup_held", 40: "think.uses_right_way", 41: "com.directions_gesture", 42: "soc.knows_familiar",
    43: "soc.secure_base", 44: "hand.stack_2", 45: "com.one_step", 46: "soc.copies_children",
    47: "", 48: "auto.tries_spoon", 49: "hand.pages_one", 50: "soc.affection",
    51: "com.body_parts", 52: "mob.pull_toy", 53: "mob.run", 54: "hand.scribble",
    55: "com.two_words", 56: "hand.tower_4", 57: "mob.kick_ball", 58: "soc.defiance",
    59: "soc.plays_next_to", 60: "auto.toothbrush", 61: "com.understands_no", 62: "",
    63: "com.names_in_book", 64: "math.sort_shape_colour", 65: "soc.pretend_objects", 66: "com.two_step",
    67: "com.pronouns", 68: "auto.toilet_interest", 69: "soc.pretend_objects", 70: "hand.tower_6",
    71: "hand.circle", 72: "mob.jump_both", 74: "com.name_age_gender", 75: "soc.notices_hurt",
    76: "com.time_words", 77: "", 78: "hand.puzzle", 79: "auto.undress",
    80: "com.action_word", 81: "think.one_colour", 82: "soc.joins_play", 83: "hand.line",
    84: "mob.throw_overhand", 85: "mob.tricycle", 86: "com.prepositions", 87: "math.big_small",
    88: "soc.mirror", 89: "soc.takes_turns", 90: "soc.pretend_role", 91: "math.count_3",
    92: "com.intelligible", 93: "mob.one_foot_2s", 94: "hand.scissors", 95: "soc.humour",
    96: "soc.pretend_objects", 97: "com.song_words", 98: "hand.circle", 99: "soc.takes_turns",
    100: "com.body_parts", 101: "com.points_book", 102: "lit.signs_logos", 103: "mob.tricycle",
    104: "math.sort_shape_colour", 105: "com.pronouns", 106: "com.prepositions", 107: "soc.follows_rules",
    108: "", 109: "math.count_10", 110: "", 111: "hand.tower_6",
    112: "com.two_step", 113: "think.names_colours", 114: "com.wh_questions", 115: "soc.pretend_role",
    116: "math.names_numbers_5", 117: "soc.joins_play", 118: "com.sentences_4", 119: "mob.one_foot_2s",
    120: "", 121: "com.prepositions", 122: "com.name_age_gender", 123: "mob.tricycle",
    124: "hand.circle", 125: "soc.pretend_role", 126: "com.three_step", 127: "hand.letters_name",
    128: "soc.follows_rules", 129: "soc.facial_expressions", 130: "auto.fork", 131: "auto.dress_some",
    132: "", 133: "think.toy_switches", 134: "", 135: "com.says_address",
    136: "auto.chores", 137: "soc.humour", 138: "soc.avoids_danger",
}
NOTES = {
    19: "Recorded at 3 months: below the WHO 1st percentile (3.8 m). Please check it was sitting WITHOUT support.",
    27: "Recorded at 5 months: below the WHO 1st percentile (5.2 m). Please check it was hands-and-knees crawling.",
    47: "No scientific equivalent kept (walking backwards).",
    62: "Too vague to observe; kept in the journal instead.",
    77: "Too vague to observe; kept in the journal instead.",
    108: "Reciting the alphabet depends on exposure, not a developmental norm; logged in the journal.",
    110: "No scientific equivalent kept (names friends).",
    120: "Covered by sentence milestones; not scored separately.",
    132: "Reading words is a later learning goal (lit.fr.read_words); not imported from a blank row.",
    134: "Too vague to observe.",
    116: "Mapped to naming numbers 1–5 (CDC 5 y).",
}


def month_to_date(m: float) -> str:
    whole = int(m)
    y, mo = BIRTH.year + (BIRTH.month - 1 + whole) // 12, (BIRTH.month - 1 + whole) % 12 + 1
    return (date(y, mo, BIRTH.day) + timedelta(days=round((m - whole) * 30.4375))).isoformat()


def main() -> None:
    ws = openpyxl.load_workbook(ROOT / "BabyPerfMGT.xlsx")["Raw ability"]
    catalog_ids = {
        row["id"] for row in csv.DictReader(open(ROOT / "catalog/early/skills.tsv"), delimiter="\t")
    }
    rows, best = [], {}
    for r in range(3, ws.max_row + 1):
        cell = ws.cell(r, 2).value
        if not cell:
            continue
        band, ability = [p.strip('"') for p in cell.split('","')][:2]
        month = ws.cell(r, 13).value
        skill = MAPPING.get(r, "")
        if skill and skill not in catalog_ids:
            raise SystemExit(f"Row {r}: unknown skill {skill}")
        rows.append([r, band, ability, "" if month is None else int(month), skill, NOTES.get(r, "")])
        if skill and month is not None:
            best[skill] = min(best.get(skill, month), month)

    out = ROOT / "catalog/seed"
    out.mkdir(parents=True, exist_ok=True)
    with open(out / "excel-mapping.tsv", "w", newline="") as f:
        w = csv.writer(f, delimiter="\t", lineterminator="\n")
        w.writerow(["excel_row", "age_band", "excel_ability", "constance_month", "skill_id", "note"])
        w.writerows(rows)
    observations = [
        {"skillId": s, "status": "achieved", "observedOn": month_to_date(m), "approximate": True,
         "note": "Imported from BabyPerfMGT.xlsx (month precision)"}
        for s, m in sorted(best.items(), key=lambda kv: (kv[1], kv[0]))
    ]
    seed = {"child": {"firstName": "Constance", "birthDate": BIRTH.isoformat(), "gestationalWeeks": None,
                      "homeLanguages": ["fr", "en", "pt"]}, "observations": observations}
    (out / "constance.json").write_text(json.dumps(seed, indent=1, ensure_ascii=False))
    print(f"{len(rows)} Excel rows, {sum(1 for r in rows if r[4])} mapped, {len(observations)} observations")


if __name__ == "__main__":
    main()
